import copy
import openpyxl
from typing import List

import concentration.fuzzy.grader as grader
from concentration.fuzzy.classes import Grade


def get_test_grades() -> List[Grade]:
    """Returns the result of grade tests.

    Grades are computed with combinations of blink rate (0 ~ 21, step size = 1)
    and body concentration value (0 ~ 1, step size = 0.1).
    """
    fuzzy_grader = grader.FuzzyGrader()
    concent_grades: List[Grade] = []
    for blink_rate in range(22):
        for body_concent in (round(0.1 * i, 1) for i in range(11)):
            grade: float = fuzzy_grader.compute_grade(blink_rate, body_concent)
            concent_grades.append(Grade(blink=blink_rate, body=body_concent, grade=grade))
    return concent_grades


def save_grades_to_spreadsheet(filename: str, grades: List[Grade]) -> None:
    """Saves the grades to the spreadsheet with titles.

    Arguments:
        filename: The spreadsheet to save grades.
        grades: The grades to save.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    # the first 0 is just a filler, also not part of the title of body
    title_of_blink: List[int] = [0] + [i for i in range(22)]
    titles_of_body: List[List[float]] = [
        [i/10] for i in range(11)
    ]
    sheet.append(title_of_blink)
    for title in titles_of_body:
        sheet.append(title)

    for grade in grades:
        # blink won't be None in the test grades
        row = chr(grade.blink + 66)  # type: ignore
        col = str(int(grade.body * 10 + 2))
        sheet[row + col] = grade.grade

    workbook.save(filename=filename)


def draw_chart_of_grades_on_spreadsheet(filename: str) -> None:
    """Draws the surface chart of grades read from the spreadsheet and save it
    back to the spreadsheet.

    Arguments:
        filename: The spreadsheet which contains the grades.
    """
    workbook = openpyxl.load_workbook(filename=filename)
    sheet = workbook.active

    # Contour chart
    contour = openpyxl.chart.SurfaceChart()
    data = openpyxl.chart.Reference(sheet, min_col=2, max_col=23, min_row=1, max_row=12)
    labels = openpyxl.chart.Reference(sheet, min_col=1, min_row=2, max_row=12)
    contour.add_data(data, titles_from_data=True)
    contour.set_categories(labels)
    contour.title = "Contour"
    contour.height = 15
    contour.width = 25
    sheet.add_chart(contour, "A15")

    # 2D Wireframe chart
    wire_frame = copy.deepcopy(contour)
    wire_frame.wireframe = True
    wire_frame.title = "2D Wireframe"
    sheet.add_chart(wire_frame, "R15")

    # Surface chart
    surface = openpyxl.chart.SurfaceChart3D()
    surface.add_data(data, titles_from_data=True)
    surface.set_categories(labels)
    surface.title = "Surface"
    surface.height = 15
    surface.width = 25
    sheet.add_chart(surface, "A45")

    # 3D Wireframe chart
    wire_frame_3d = copy.deepcopy(surface)
    wire_frame_3d.wireframe = True
    wire_frame_3d.title = "3D Wireframe"
    sheet.add_chart(wire_frame_3d, "R45")

    workbook.save(filename=filename)


if __name__ == "__main__":
    spreadsheet: str = "test_grades.xlsx"

    grades: List[Grade] = get_test_grades()
    save_grades_to_spreadsheet(spreadsheet, grades)
    draw_chart_of_grades_on_spreadsheet(spreadsheet)

    # the following is to demo the parsing with json
    import concentration.fuzzy.parse as parse

    json_file: str = "test_grades.json"

    parse.save_grades_to_json(json_file, grades)
    parsed_grades: List[Grade] = parse.read_grades_from_json(json_file)
    # test correctness
    assert parsed_grades == grades
