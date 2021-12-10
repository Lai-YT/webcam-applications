import copy
import openpyxl
from typing import List, Optional, Tuple, Union


import fuzzy.grader as grader
import fuzzy.parse as parse


def output_grade_text(filename: str) -> None:
    """Outputs the data texts of grade in to file.

    Grades are computed with combinations of blink rate (0 ~ 21, step size = 1)
    and body concentration value (0 ~ 1, step size = 0.1).
    Arguments:
        filename: The file to output the grades to. It's opened with "w+" mode.
    """
    fuzzy_grader = grader.FuzzyGrader()
    # grade, blink rate, body concent
    concent_grades: List[Tuple[float, int, float]] = []
    for blink_rate in range(22):
        for body_concent in (round(0.1 * i, 1) for i in range(11)):
            grade: float = fuzzy_grader.compute_grade(blink_rate, body_concent)
            concent_grades.append((grade, blink_rate, body_concent))

    with open(filename, mode="w+") as f:
        for grade, blink, body in concent_grades:
            f.write(f"{blink}\n"
                    + f"{body}\n"
                    + f"{grade}\n"
                    + "---\n")


def output_grade_spreadsheet(text_file: str, spreadsheet_file: str) -> None:
    """Outputs the data texts into the spreadsheet.

    Arguments:
        text_file: The file which contains the data text.
        spreadsheet_file: The file to output the spreadsheet to.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    grades = parse.parse_grades(text_file)

    titles: List[List[Union[int, float, None]]] = [
        [None] + [i for i in range(22)],
    ]
    for i in range(11):
        titles.append([i/10])

    for row in titles:
        sheet.append(row)

    for grade in grades:
        row: str = chr(grade.blink + 66)
        col = str(int(grade.body * 10 + 2))
        sheet[row + col] = grade.grade

    workbook.save(filename=spreadsheet_file)


def output_grade_chart(
        spreadsheet_file: str,
        chart_file: Optional[str] = None) -> None:
    """Outputs the charts of the data held by spreadsheet.

    Arguments:
        spreadsheet_file: The spreadsheet to read the data from.
        chart_file:
            The file to output the charts to. The charts are output to the
            spreadsheet_file if chart_file isn't given.
    """
    if chart_file is None:
        chart_file = spreadsheet_file

    workbook = openpyxl.load_workbook(filename=spreadsheet_file)
    sheet = workbook.active

    # Contour chart
    contour = openpyxl.chart.SurfaceChart()
    data = openpyxl.chart.Reference(sheet, min_col=2, max_col=23, min_row=1, max_row=12)
    labels = openpyxl.chart.Reference(sheet, min_col=1, min_row=2, max_row=12)
    contour.add_data(data, titles_from_data=True)
    contour.set_categories(labels)
    contour.title = "Contour"
    sheet.add_chart(contour, "A15")

    # 2D Wireframe chart
    wire_frame = copy.deepcopy(contour)
    wire_frame.wireframe = True
    wire_frame.title = "2D Wireframe"
    sheet.add_chart(wire_frame, "L15")

    # Surface chart
    surface = openpyxl.chart.SurfaceChart3D()
    surface.add_data(data, titles_from_data=True)
    surface.set_categories(labels)
    surface.title = "Surface"
    sheet.add_chart(surface, "A31")

    # 3D Wireframe chart
    wire_frame_3d = copy.deepcopy(surface)
    wire_frame_3d.wireframe = True
    wire_frame_3d.title = "3D Wireframe"
    sheet.add_chart(wire_frame_3d, "L31")

    workbook.save(filename=chart_file)


if __name__ == "__main__":
    output_grade_text("fuzzy_grades.txt")
    output_grade_spreadsheet("fuzzy_grades.txt", "fuzzy_grades.xlsx")
    output_grade_chart("fuzzy_grades.xlsx")
